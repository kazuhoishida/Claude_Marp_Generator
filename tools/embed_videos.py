#!/usr/bin/env python3
"""marp が書き出した pptx を、配布できる形に整える後処理。

やること:
  1. 白いだけの全面シェイプ（変換の副産物）を取り除く
  2. 折り返しで複数に割れたテキストボックスを、段落ごとに1つへまとめる
  3. テキストボックスの左右の余白を対称にする
  4. 指定したスライドに動画を貼り込む（PDF には動画が入らないので pptx だけ）

    python3 tools/embed_videos.py <pptx> [--assets DIR] [--map videos.json]

対応表（既定では素材ディレクトリの videos.json）はこの形で書く。
キーはスライド番号（1始まり）、値は [動画, ポスターに使う静止画]。

    {
      "4": ["IMG_3082.mp4", "01_float_two.png"],
      "5": ["IMG_3086.mp4", "02_rotate_styrofoam.png"]
    }

必要なもの: python-pptx（`pip install python-pptx`）と ffmpeg。
"""

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt

SLIDE_W = 1280           # png の書き出し幅
SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
EMU_PER_PX = 9525        # 12192000 EMU / 1280 px
SLIDE_CENTER = SLIDE_W_EMU // 2
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MARGIN = 0.06            # テキストボックスに持たせる余裕
LATIN = re.compile(r"[0-9A-Za-z）)\]]$")
LATIN_HEAD = re.compile(r"^[0-9A-Za-z（(\[]")


# --- 背景 -----------------------------------------------------------------

def drop_blank_background(prs):
    """白いだけの全面シェイプを消す。

    PDF 経由の変換では、どのスライドにもページの地を描いた全面シェイプが
    3枚重なる。白いスライドでは白い板が3枚乗っているだけで、選択や編集の
    邪魔になる。色が付いているもの（章扉のグレーなど）は地の色そのものなので
    残す。
    """
    removed = 0
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if not _is_full_slide(shape):
                continue
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                continue
            if _fill_color(shape) == "FFFFFF":
                shape._element.getparent().remove(shape._element)
                removed += 1
    return removed


def _fill_color(shape):
    """シェイプの塗り色（RRGGBB）。塗りが無ければ None。

    文字色などを拾わないよう、spPr の直下の solidFill だけを見る。
    """
    spPr = shape._element.find(f"{{{NS_P}}}spPr")
    if spPr is None:
        return None
    fill = spPr.find(f"{{{NS_A}}}solidFill")
    if fill is None:
        return None
    clr = fill.find(f"{{{NS_A}}}srgbClr")
    return clr.get("val", "").upper() if clr is not None else None


def _is_full_slide(shape):
    return shape.width >= SLIDE_W_EMU * 0.98 and shape.height >= SLIDE_H_EMU * 0.98


# --- テキスト -------------------------------------------------------------

def _lines(slide):
    """テキストの入ったシェイプを、上から順に返す。"""
    out = [sh for sh in slide.shapes
           if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
    return sorted(out, key=lambda sh: (sh.top, sh.left))


def _font_size(shape):
    """1つ目の run の文字サイズ（pt）。"""
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return None


def _bold(shape):
    """1つ目の run が太字かどうか。"""
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            return bool(run.font.bold)
    return False


def _same_paragraph(a, b):
    """b が a の続きの行かどうか。

    左端・文字サイズ・太さが同じで、行送りが文字サイズの2倍以内なら
    同じ段落とみなす。箇条書きの項目同士は li の余白があるぶん行送りが
    広くなるので分かれる。小見出し（h3）は本文と同じ文字サイズなので、
    太さを見ないと直下の本文とつながってしまう。
    """
    sa, sb = _font_size(a), _font_size(b)
    if sa is None or sb is None or abs(sa - sb) > 0.5:
        return False
    if _bold(a) != _bold(b):
        return False
    if abs(a.left - b.left) > 20000:            # 約2px
        return False
    pitch_pt = (b.top - a.top) / 12700          # EMU -> pt
    return 1.1 * sa < pitch_pt < 2.0 * sa


def _merge(slide, group):
    """group（同じ段落の行）を先頭のシェイプに統合し、残りを削除する。"""
    head = group[0]
    size_pt = _font_size(head) or 16
    pitch_pt = (group[1].top - head.top) / 12700

    text = ""
    for sh in group:
        line = sh.text_frame.text.strip()
        if text and LATIN.search(text) and LATIN_HEAD.match(line):
            text += " "          # 欧文どうしは単語がつながらないよう空白を入れる
        text += line

    para = head.text_frame.paragraphs[0]
    for run in para.runs[1:]:
        run._r.getparent().remove(run._r)
    para.runs[0].text = text
    # 行送りは実寸で指定する。倍数で指定するとフォント既定の行高（およそ
    # 文字サイズの1.2倍）に対する倍率になり、元より広くなる。
    para.line_spacing = Pt(pitch_pt)

    head.width = max(sh.width for sh in group) + int(head.width * 0.02)
    head.height = group[-1].top + group[-1].height - head.top
    head.text_frame.word_wrap = True

    for sh in group[1:]:
        sh._element.getparent().remove(sh._element)


def merge_lines(prs):
    """折り返しで割れた行を段落ごとに1つのテキストボックスへまとめる。

    --pptx-editable は PDF 経由の変換なので、1行 = 1テキストボックスになる。
    そのままだと PowerPoint で文言を直しても行が流れ直さず、改行位置が
    固定されたままになる。
    """
    merged = 0
    for slide in prs.slides:
        group = []
        for shape in _lines(slide) + [None]:
            if group and (shape is None or not _same_paragraph(group[-1], shape)):
                if len(group) > 1:
                    _merge(slide, group)
                    merged += 1
                group = []
            if shape is not None:
                group.append(shape)
    return merged


def unwrap_text(prs):
    """1行だけの箱は折り返しを切り、幅に余裕を持たせる。

    箱の幅は元の文字幅ぴったりに作られるため、開く側の組版がわずかに広いと
    最後の1文字が折り返して2行になり、下の要素と重なる。中央に置かれた箱は
    左右均等に広げないと中心がずれる。
    """
    n = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if not shape.text_frame.text.strip():
                continue
            if shape.text_frame.word_wrap:      # まとめた段落は折り返しを残す
                continue
            shape.text_frame.word_wrap = False
            grow = int(shape.width * MARGIN)
            centered = abs(shape.left + shape.width // 2 - SLIDE_CENTER) < 40000
            shape.width += grow
            if centered:
                shape.left -= grow // 2
            n += 1
    return n


def _overlaps(a, others):
    """a の矩形が他のシェイプと重なるか。"""
    for b in others:
        if b is a or _is_full_slide(b):      # 地の全面シェイプは無視する
            continue
        if (a.left < b.left + b.width and a.left + a.width > b.left
                and a.top < b.top + b.height and a.top + a.height > b.top):
            return True
    return False


def balance_width(prs):
    """本文のテキストボックスを、左右の余白が同じになる幅にそろえる。

    PDF 経由の変換では、箱の幅が「その行の文字が占める幅」になるため、
    右端が行ごとにばらつき、スライドに対して左右非対称に見える。
    左の余白と同じだけ右にも余白が残る幅に広げてそろえる。

    ヘッダーやページ番号（上下の端）、右半分から始まる箱（多カラムの
    右側など）は対象にしない。広げると他の要素に重なる場合も見送る。
    """
    n = 0
    for slide in prs.slides:
        shapes = list(slide.shapes)
        for shape in shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if not shape.text_frame.text.strip():
                continue
            if shape.left >= SLIDE_CENTER:                      # 右半分から始まる箱
                continue
            if shape.top < SLIDE_H_EMU * 0.08 or shape.top > SLIDE_H_EMU * 0.9:
                continue                                        # ヘッダー / ページ番号
            target = SLIDE_W_EMU - 2 * shape.left
            if target <= shape.width:
                continue
            before = shape.width
            shape.width = target
            if _overlaps(shape, shapes):
                shape.width = before                            # 重なるなら戻す
                continue
            n += 1
    return n


# --- 動画 -----------------------------------------------------------------

def video_size(path):
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-select_streams", "v:0",
         "-show_entries", "stream=width,height", "-of", "csv=p=0:s=x", str(path)],
        capture_output=True, text=True, check=True).stdout.strip()
    w, h = out.split("x")[:2]
    return int(w), int(h)


def box_from_picture(slide):
    """スライド上でいちばん大きい画像シェイプの矩形（EMU）。無ければ None。"""
    pics = [s for s in slide.shapes if s.shape_type == 13]  # PICTURE
    if not pics:
        return None
    p = max(pics, key=lambda s: s.width * s.height)
    return p.left, p.top, p.width, p.height


def box_from_png(png_path):
    """スライドの png から図版の矩形を実測する（EMU）。

    --pptx（各ページが1枚の背景画像）のときはシェイプが無いのでこちらを使う。
    """
    im = Image.open(png_path).convert("L")
    w, h = im.size
    px = im.load()
    x0, y0, x1, y1 = w, h, 0, 0
    for y in range(int(h * 0.19), int(h * 0.94)):
        for x in range(w):
            if px[x, y] < 235:      # 白地ではない
                x0, y0 = min(x0, x), min(y0, y)
                x1, y1 = max(x1, x), max(y1, y)
    if x1 <= x0 or y1 <= y0:
        raise SystemExit(f"図版が見つからない: {png_path}")
    k = SLIDE_W / w * EMU_PER_PX
    return [round(v * k) for v in (x0, y0, x1 - x0 + 1, y1 - y0 + 1)]


def fit(box, ratio):
    """矩形 box の内側に、縦横比 ratio (w/h) を保って収めた矩形を返す。

    枠の比率に合わせて引き伸ばすと、縦位置で撮った動画が横に伸びる。
    """
    x, y, w, h = box
    if w / h > ratio:
        nw, nh = round(h * ratio), h
    else:
        nw, nh = w, round(w / ratio)
    return x + (w - nw) // 2, y + (h - nh) // 2, nw, nh


def main():
    ap = argparse.ArgumentParser(description="pptx に動画を貼り込み、段落を整える")
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--assets", type=Path, help="素材ディレクトリ（既定: output/assets/<デッキ名>）")
    ap.add_argument("--map", dest="mapping", type=Path, help="対応表 json（既定: <素材>/videos.json）")
    args = ap.parse_args()

    assets = args.assets or Path("output/assets") / args.pptx.stem
    mapping = args.mapping or assets / "videos.json"
    if not mapping.exists():
        sys.exit(f"対応表が無い: {mapping}")
    targets = json.loads(mapping.read_text())

    prs = Presentation(args.pptx)
    print(f"白い全面シェイプを削除: {drop_blank_background(prs)} 枚")
    print(f"段落をまとめた: {merge_lines(prs)} 箇所")
    print(f"折り返しを無効化: {unwrap_text(prs)} 個のテキストボックス")
    print(f"左右の余白をそろえた: {balance_width(prs)} 個のテキストボックス")

    for no, (movie, poster) in sorted(targets.items(), key=lambda kv: int(kv[0])):
        no = int(no)
        slide = prs.slides[no - 1]
        box = box_from_picture(slide)
        source = "picture"
        if box is None:
            box = box_from_png(args.pptx.with_suffix("").parent /
                               f"{args.pptx.stem}.{no:03d}.png")
            source = "png"

        vw, vh = video_size(assets / movie)
        x, y, w, h = fit(box, vw / vh)
        slide.shapes.add_movie(
            str(assets / movie), Emu(x), Emu(y), Emu(w), Emu(h),
            poster_frame_image=str(assets / poster),
            mime_type="video/mp4",
        )
        print(f"slide {no:2d}: {movie:14s} {vw}x{vh}"
              f" -> ({x // EMU_PER_PX},{y // EMU_PER_PX})"
              f" {w // EMU_PER_PX}x{h // EMU_PER_PX}px  [{source}]")

    prs.save(args.pptx)
    print(f"saved {args.pptx} ({args.pptx.stat().st_size / 1_000_000:.1f}MB)")


if __name__ == "__main__":
    sys.exit(main())
